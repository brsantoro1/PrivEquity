import openpyxl
from openpyxl.utils.cell import range_boundaries
from datetime import datetime, timedelta

'This is a new line designed to test our git push'

def get_column_headers(workbook_path, sheet_name, table_name):
    """
    Returns a dictionary mapping header names to column indices for a specified table in an Excel workbook.
    
    Args:
    workbook_path (str): The file path to the Excel workbook.
    sheet_name (str): The name of the sheet in the workbook.
    table_name (str): The name of the table in the sheet.

    Returns:
    dict: A dictionary mapping header names to their column indices.
    """

    def excel_date_to_datetime(excel_date):
        return datetime(1899, 12, 30) + timedelta(days=excel_date)

    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    sheet = wb[sheet_name]
    table = next((tbl for tbl in sheet.tables.values() if tbl.name == table_name), None)

    if not table:
        wb.close()
        raise ValueError(f"Table '{table_name}' not found in sheet '{sheet_name}'")

    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    headers = {}

    for col_index, col in enumerate(sheet.iter_cols(min_col=min_col, max_col=max_col, min_row=min_row, max_row=min_row), start=1):
        header_cell = col[0]
        headers[header_cell.value] = col_index - 1

    wb.close()
    return headers

# Example usage
#table_headers = get_column_headers("path_to_your_workbook.xlsx", "Sheet Name", "Table Name")
#print(p_table_headers)

#table_headers = get_column_headers(r"K:\Market Maps\Interest Rates Map.xlsm", "Master", "Master")
#print(table_headers)